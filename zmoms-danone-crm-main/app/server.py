import http.server
import socketserver
import json
import os
import urllib.parse
import datetime

PORT = int(os.environ.get('PORT', 8080))
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, 'data', 'crm_store.json')
INITIAL_DATA_FILE = os.path.join(BASE_DIR, 'data', 'initial_data.js')
EXCEL_FILE = os.path.join(BASE_DIR, 'MASTER_DATA_INTERNAL_v2.xlsx')

def load_db():
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    return {}

def save_db(db):
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(db, f, ensure_ascii=False, indent=2)
    try:
        with open(INITIAL_DATA_FILE, 'w', encoding='utf-8') as f:
            f.write('window.CRM_INITIAL_DATA = ' + json.dumps(db, ensure_ascii=False, indent=2) + ';\n')
    except Exception as e:
        print(f"Error saving initial_data.js: {e}")

def append_order_to_excel(order):
    try:
        if not os.path.exists(EXCEL_FILE):
            return
        import openpyxl
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if 'ORDERS' not in wb.sheetnames:
            return
        ws = wb['ORDERS']
        
        # Tìm dòng trống đầu tiên
        row_idx = 2
        while ws.cell(row=row_idx, column=1).value is not None:
            row_idx += 1
            
        stt = row_idx - 1
        ws.cell(row=row_idx, column=1, value=stt)
        ws.cell(row=row_idx, column=2, value=order.get('date', datetime.datetime.now().strftime('%d/%m/%Y')))
        ws.cell(row=row_idx, column=3, value=order.get('phone', ''))
        ws.cell(row=row_idx, column=4, value=f'=IF($C{row_idx}="","",IFERROR(INDEX(DATA_MASTER!$C$2:$C$5000,MATCH($C{row_idx},DATA_MASTER!$D$2:$D$5000,0)),"{order.get("customer_name", "")}"))')
        ws.cell(row=row_idx, column=5, value=f'=IF($C{row_idx}="","",IFERROR(INDEX(DATA_MASTER!$I$2:$I$5000,MATCH($C{row_idx},DATA_MASTER!$D$2:$D$5000,0)),"1. Mang thai 3 tháng đầu"))')
        ws.cell(row=row_idx, column=6, value=f'=IF($C{row_idx}="","",IFERROR(INDEX(DATA_MASTER!$L$2:$L$5000,MATCH($C{row_idx},DATA_MASTER!$D$2:$D$5000,0)),"ZM"))')
        ws.cell(row=row_idx, column=7, value=order.get('product', 'Sữa Aptamil 380g'))
        ws.cell(row=row_idx, column=8, value=1)
        ws.cell(row=row_idx, column=9, value=320000)
        ws.cell(row=row_idx, column=10, value=f'=IF(OR($H{row_idx}="",$I{row_idx}=""),"",$H{row_idx}*$I{row_idx})')
        ws.cell(row=row_idx, column=11, value='Chuyển khoản QR')
        ws.cell(row=row_idx, column=12, value='Chờ chuyển khoản')
        ws.cell(row=row_idx, column=13, value='')
        ws.cell(row=row_idx, column=14, value=order.get('phone', ''))
        ws.cell(row=row_idx, column=15, value=order.get('customer_name', ''))
        full_addr = f"{order.get('address', '')}, {order.get('district', '')}".strip(', ')
        ws.cell(row=row_idx, column=16, value=full_addr)
        ws.cell(row=row_idx, column=17, value='Chờ giao')
        ws.cell(row=row_idx, column=18, value=order.get('shop', 'Voi Con'))
        ws.cell(row=row_idx, column=19, value=f'=IF(AND($A{row_idx}="",$C{row_idx}=""),"",IF($A{row_idx}="","Thiếu mã đơn; ","")&IF(COUNTIF($A$2:$A{row_idx},$A{row_idx})>1,"Trùng mã đơn; ","")&IF(COUNTIF($C$2:$C{row_idx},$C{row_idx})>1,"Trùng SĐT; ","")&IF(ISNA(MATCH($C{row_idx},DATA_MASTER!$D$2:$D$5000,0)),"SĐT chưa đăng ký hội viên; ","")&IF(AND($L{row_idx}="Đã chuyển khoản",$M{row_idx}=""),"Thiếu ngày nhận CK; ",""))')
        ws.cell(row=row_idx, column=20, value=f"Mã đơn: {order.get('code', '')} | Phễu: {order.get('funnel', '')}")
        
        wb.save(EXCEL_FILE)
        print(f"Appended order {order.get('code')} to Excel sheet ORDERS row {row_idx}")
    except Exception as e:
        print(f"Error appending order to Excel: {e}")

class CRMHandler(http.server.SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=BASE_DIR, **kwargs)

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        query = urllib.parse.parse_qs(parsed.query)

        if path.startswith('/api/'):
            self.handle_api_get(path, query)
        else:
            if path == '/' or path == '':
                self.path = '/index.html'
            super().do_GET()

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        content_length = int(self.headers.get('Content-Length', 0))
        post_data = self.rfile.read(content_length) if content_length > 0 else b'{}'
        try:
            body = json.loads(post_data.decode('utf-8'))
        except Exception:
            body = {}

        if path.startswith('/api/'):
            self.handle_api_post(path, body)
        else:
            self.send_error(404, "Not Found")

    def do_PUT(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        content_length = int(self.headers.get('Content-Length', 0))
        post_data = self.rfile.read(content_length) if content_length > 0 else b'{}'
        try:
            body = json.loads(post_data.decode('utf-8'))
        except Exception:
            body = {}

        if path.startswith('/api/'):
            self.handle_api_put(path, body)
        else:
            self.send_error(404, "Not Found")

    def send_json(self, data, status_code=200):
        self.send_response(status_code)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, PUT, DELETE, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()
        self.wfile.write(json.dumps(data, ensure_ascii=False).encode('utf-8'))

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, PUT, DELETE, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type')
        self.end_headers()

    def handle_api_get(self, path, query):
        db = load_db()

        if path == '/api/stats':
            orders = db.get('orders', [])
            weeks = db.get('weeks', [])
            total_orders = len(orders)
            delivered_orders = len([o for o in orders if 'Đã giao' in o.get('shop_status', '') or 'Hoàn tất' in o.get('overall_status', '')])
            pending_orders = len([o for o in orders if 'Chờ' in o.get('shop_status', '') or 'Chưa lên' in o.get('shop_status', '')])
            shipping_orders = len([o for o in orders if 'Đang giao' in o.get('shop_status', '') or 'Shop đã nhận' in o.get('shop_status', '')])
            
            sla_alert_count = len([o for o in orders if 'Chờ' in o.get('shop_status', '') or 'Chưa' in o.get('shop_status', '')])

            meta = db.get('meta', {})
            total_target_data = meta.get('target_data', 10010)
            total_target_nu = meta.get('target_nu', 2517)
            total_target_cr = meta.get('target_cr', 25.14)
            calc_mode = meta.get('calc_mode', 'auto')
            
            actual_data_sum = sum(w.get('actual_data', 0) for w in weeks)
            actual_nu_sum = sum(w.get('actual_nu', 0) for w in weeks)
            if actual_nu_sum < total_orders:
                actual_nu_sum = total_orders

            raw_leads = db.get('raw_leads', [])
            total_raw = meta.get('total_raw_data', len(raw_leads)) if calc_mode == 'manual' else (len(raw_leads) or 423)
            valid_raw = meta.get('valid_raw_data', 409) if calc_mode == 'manual' else (len([r for r in raw_leads if r.get('valid_kpi')]) or 409)
            invalid_raw = (total_raw - valid_raw) if calc_mode == 'manual' else (len([r for r in raw_leads if not r.get('valid_kpi')]) or 14)
            converted_raw = meta.get('converted_raw_data', 185) if calc_mode == 'manual' else (len([r for r in raw_leads if r.get('has_ordered')]) or 185)

            stats = {
                'project_name': meta.get('project_name', 'Đề kháng vững vàng, trí tuệ tinh anh'),
                'duration': meta.get('duration', '16 Tuần (01/09/2026 - 15/12/2026)'),
                'target_data': total_target_data,
                'target_nu': total_target_nu,
                'target_cr': total_target_cr,
                'calc_mode': calc_mode,
                'actual_data': actual_data_sum,
                'percent_data': round((actual_data_sum / total_target_data) * 100, 1) if total_target_data else 0,
                'actual_nu': actual_nu_sum,
                'percent_nu': round((actual_nu_sum / total_target_nu) * 100, 1) if total_target_nu else 0,
                'total_orders': total_orders,
                'delivered_orders': delivered_orders,
                'pending_orders': pending_orders,
                'shipping_orders': shipping_orders,
                'sla_alert_count': sla_alert_count,
                'total_raw_data': total_raw,
                'valid_raw_data': valid_raw,
                'invalid_raw_data': invalid_raw,
                'converted_raw_data': converted_raw,
                'approved_hospitals': meta.get('approved_hospitals', 11),
                'total_hospitals': meta.get('total_hospitals', 26),
                'raw_teams': meta.get('raw_teams', {}),
                'raw_channels': meta.get('raw_channels', {}),
                'total_unique_customers': len(db.get('customers', [])),
                'channels': {
                    'hospital': {'leads': 450, 'nu': 120, 'name': 'Hội thảo Bệnh viện'},
                    'online': {'leads': 580, 'nu': 142, 'name': 'Online Webinars'},
                    'spa': {'leads': 326, 'nu': 79, 'name': 'Đối tác Spa Mẹ Bầu'},
                    'midwife': {'leads': 148, 'nu': 47, 'name': 'Nữ Hộ Sinh Tại Nhà'}
                }
            }
            self.send_json({'success': True, 'data': stats})

        elif path == '/api/raw-leads':
            raw_leads = db.get('raw_leads', [])
            q_search = query.get('q', [None])[0]
            team_filter = query.get('team', [None])[0]
            valid_filter = query.get('valid', [None])[0]
            ordered_filter = query.get('ordered', [None])[0]

            filtered = raw_leads
            if team_filter:
                filtered = [r for r in filtered if r.get('team', '').lower() == team_filter.lower()]
            if valid_filter is not None:
                is_v = valid_filter.lower() in ['true', '1', 'có']
                filtered = [r for r in filtered if r.get('valid_kpi') == is_v]
            if ordered_filter is not None:
                is_o = ordered_filter.lower() in ['true', '1', 'có']
                filtered = [r for r in filtered if r.get('has_ordered') == is_o]
            if q_search:
                q = q_search.lower()
                filtered = [r for r in filtered if q in r.get('name', '').lower() or q in r.get('phone', '').lower() or q in r.get('province', '').lower() or q in r.get('pancake_id', '').lower() or q in r.get('ref_code', '').lower() or q in r.get('source', '').lower()]

            self.send_json({'success': True, 'data': filtered, 'total': len(filtered), 'meta': db.get('meta', {})})

        elif path == '/api/orders':
            orders = db.get('orders', [])
            shop_filter = query.get('shop', [None])[0]
            status_filter = query.get('status', [None])[0]
            q_search = query.get('q', [None])[0]

            filtered = orders
            if shop_filter:
                filtered = [o for o in filtered if o.get('shop', '').lower() == shop_filter.lower()]
            if status_filter:
                filtered = [o for o in filtered if status_filter.lower() in o.get('shop_status', '').lower()]
            if q_search:
                q = q_search.lower()
                filtered = [o for o in filtered if q in o.get('code', '').lower() or q in o.get('customer_name', '').lower() or q in o.get('phone', '').lower() or q in o.get('district', '').lower()]

            self.send_json({'success': True, 'data': filtered, 'total': len(filtered)})

        elif path == '/api/customers':
            customers = db.get('customers', [])
            q_search = query.get('q', [None])[0]
            stage_filter = query.get('stage', [None])[0]
            channel_filter = query.get('channel', [None])[0]

            filtered = customers
            if stage_filter:
                filtered = [c for c in filtered if stage_filter.lower() in c.get('stage', '').lower()]
            if channel_filter:
                filtered = [c for c in filtered if channel_filter.lower() in c.get('channel', '').lower()]
            if q_search:
                q = q_search.lower()
                filtered = [c for c in filtered if q in c.get('id', '').lower() or q in c.get('name', '').lower() or q in c.get('phone', '').lower() or q in c.get('district', '').lower()]

            self.send_json({'success': True, 'data': filtered, 'total': len(filtered)})

        elif path == '/api/shops/summary' or path == '/api/shops':
            shops = db.get('mbi_shops', [])
            orders = db.get('orders', [])
            
            summary = []
            for s in shops:
                s_name = s.get('name', '')
                shop_orders = [o for o in orders if o.get('shop', '').strip().lower() == s_name.strip().lower()]
                
                total = len(shop_orders)
                delivered = len([o for o in shop_orders if 'Đã giao' in o.get('shop_status', '')])
                shipping = len([o for o in shop_orders if 'Đang giao' in o.get('shop_status', '')])
                pending = len([o for o in shop_orders if 'Chờ' in o.get('shop_status', '') or 'Chưa' in o.get('shop_status', '') or 'Shop đã nhận' in o.get('shop_status', '')])
                quota = s.get('quota', 100)
                completion_rate = round((delivered / total * 100), 1) if total > 0 else 0
                remaining = max(0, quota - total)

                summary.append({
                    **s,
                    'total_orders': total,
                    'delivered_orders': delivered,
                    'shipping_orders': shipping,
                    'pending_orders': pending,
                    'remaining_quota': remaining,
                    'completion_rate': completion_rate
                })

            self.send_json({'success': True, 'data': summary})

        elif path == '/api/hospitals':
            self.send_json({'success': True, 'data': db.get('hospitals', [])})

        elif path == '/api/midwives':
            self.send_json({'success': True, 'data': db.get('midwives', [])})

        elif path == '/api/webinars':
            self.send_json({'success': True, 'data': db.get('webinars', [])})

        elif path == '/api/spas':
            self.send_json({'success': True, 'data': db.get('spas', [])})

        elif path == '/api/weeks':
            self.send_json({'success': True, 'data': db.get('weeks', [])})

        elif path == '/api/action-items':
            self.send_json({'success': True, 'data': db.get('action_items', [])})

        elif path == '/api/kpi':
            team = db.get('sales_team', [])
            results = []
            for m in team:
                target = m.get('target_nu', 100)
                actual = m.get('actual_nu', 0)
                subsidized = m.get('subsidized_nu', 0)
                base_rev = m.get('base_revenue', 0)
                is_cdb = m.get('team') == 'CDB'
                team_multiplier = 0.7 if is_cdb else 1.0

                pct = (actual / target) if target > 0 else 0

                # 1. Thưởng NU bán
                rate_nu = 0
                if pct >= 1.10: rate_nu = 30000
                elif pct >= 0.95: rate_nu = 25000
                elif pct >= 0.90: rate_nu = 20000
                bonus_nu = actual * rate_nu * team_multiplier

                # 2. Thưởng NU trợ giá
                pct_sub = (subsidized / (target * 0.2)) if target > 0 else 0
                rate_sub = 0
                if pct_sub >= 0.91: rate_sub = 5000
                elif pct_sub >= 0.80: rate_sub = 3000
                bonus_sub = subsidized * rate_sub * team_multiplier

                # 3. Thưởng Doanh số
                rate_rev = 0
                if pct >= 0.90: rate_rev = 0.03
                elif pct >= 0.80: rate_rev = 0.02
                bonus_rev = base_rev * rate_rev * team_multiplier

                total_bonus = bonus_nu + bonus_sub + bonus_rev

                results.append({
                    **m,
                    'achievement_rate': round(pct * 100, 1),
                    'bonus_nu': bonus_nu,
                    'bonus_sub': bonus_sub,
                    'bonus_rev': bonus_rev,
                    'total_bonus': total_bonus,
                    'team_multiplier': team_multiplier
                })

            self.send_json({'success': True, 'data': results})

        elif path == '/api/export':
            orders = db.get('orders', [])
            csv_lines = ['Mã Đơn,Ngày Đặt,Tên Mẹ Bầu,Số Điện Thoại,Khu Vực,Sản Phẩm,Shop MBI,Trạng Thái Shop,Đơn Vị Vận Chuyển,Mã Vận Đơn,Trạng Thái Tổng Thể']
            for o in orders:
                line = f'"{o.get("code","")}","{o.get("date","")}","{o.get("customer_name","")}","{o.get("phone","")}","{o.get("district","")}","{o.get("product","")}","{o.get("shop","")}","{o.get("shop_status","")}","{o.get("carrier","")}","{o.get("tracking_code","")}","{o.get("overall_status","")}"'
                csv_lines.append(line)
            
            csv_data = '\r\n'.join(csv_lines).encode('utf-8-sig')
            self.send_response(200)
            self.send_header('Content-Type', 'text/csv; charset=utf-8')
            self.send_header('Content-Disposition', 'attachment; filename="Danone_Zmoms_Orders_Export.csv"')
            self.end_headers()
            self.wfile.write(csv_data)

        else:
            self.send_error(404, "Endpoint not found")

    def handle_api_post(self, path, body):
        db = load_db()

        if path == '/api/orders/check-duplicate':
            phone = str(body.get('phone', '')).replace(' ', '').replace('.', '').replace('-', '').strip()
            if phone.startswith('84'): phone = '0' + phone[2:]
            elif not phone.startswith('0') and len(phone) in [9, 10]: phone = '0' + phone

            orders = db.get('orders', [])
            match = None
            for o in orders:
                op = str(o.get('phone', '')).replace(' ', '').replace('.', '').strip()
                if op == phone and phone:
                    match = o
                    break

            if match:
                self.send_json({
                    'is_duplicate': True,
                    'match': match,
                    'message': f"⚠️ CẢNH BÁO TRÙNG LẶP: SĐT {phone} đã có đơn hàng {match.get('code')} ngày {match.get('date')} ({match.get('customer_name')}) tại shop {match.get('shop')}!"
                })
            else:
                self.send_json({
                    'is_duplicate': False,
                    'message': f"✅ SĐT {phone} hợp lệ! Chưa có đơn hàng trên hệ thống."
                })

        elif path == '/api/orders/auto-dispatch':
            region_input = str(body.get('region', '')).strip().lower()
            district_input = str(body.get('district', '')).strip().lower()
            address_input = str(body.get('address', '')).strip().lower()
            pref_shop = str(body.get('preferred_shop', '')).strip()
            amount = body.get('amount', 320000)

            full_text = f"{region_input} {district_input} {address_input}".lower()
            shops = db.get('mbi_shops', [])

            # 1. Định nghĩa cụm vùng và shop tương ứng theo yêu cầu:
            # - Cần Thơ: Thế Giới Sơ Sinh, Nous Baby, Nhà Của Gạo
            # - Hồ Chí Minh: Voi Con
            # - Hà Nội: Mẹ Ơi, Sữa 3 Tốt
            # - Hải Phòng: Ababy, Cute Baby
            # - Khác: Long Châu
            can_tho_kw = ['cần thơ', 'can tho', 'ninh kiều', 'cái răng', 'bình thủy', 'ô môn', 'thốt nốt', 'phong điền', 'thới lai', 'cờ đỏ', 'vĩnh thạnh']
            hcm_kw = ['hồ chí minh', 'ho chi minh', 'hcm', 'sài gòn', 'sai gon', 'tp.hcm', 'tphcm', 'quận 1', 'quận 2', 'quận 3', 'quận 4', 'quận 5', 'quận 6', 'quận 7', 'quận 8', 'quận 9', 'quận 10', 'quận 11', 'quận 12', 'q1', 'q2', 'q3', 'q7', 'q12', 'thủ đức', 'bình thạnh', 'gò vấp', 'tân bình', 'tân phú', 'phú nhuận', 'bình tân', 'hóc môn', 'củ chi', 'bình chánh', 'nhà bè', 'cần giờ']
            hn_kw = ['hà nội', 'ha noi', 'hn', 'ba đình', 'hoàn kiếm', 'tây hồ', 'long biên', 'cầu giấy', 'đống đa', 'hai bà trưng', 'hoàng mai', 'thanh xuân', 'sóc sơn', 'đông anh', 'gia lâm', 'nam từ liêm', 'thanh trì', 'bắc từ liêm', 'mê linh', 'hà đông', 'sơn tây', 'ba vì', 'phúc thọ', 'đan phượng', 'hoài đức', 'quốc oai', 'thạch thất', 'chương mỹ', 'thanh oai', 'thường tín', 'phú xuyên', 'ứng hòa', 'mỹ đức']
            hp_kw = ['hải phòng', 'hai phong', 'hp', 'hồng bàng', 'ngô quyền', 'lê chân', 'hải an', 'kiến an', 'đồ sơn', 'dương kinh', 'thủy nguyên', 'an dương', 'an lão', 'kiến thụy', 'tiên lãng', 'vĩnh bảo', 'cát hải', 'bạch long vĩ']

            region_code = 'other'
            region_name = 'Khu vực khác (Toàn quốc)'
            target_shop_names = ['Long Châu']

            if 'can_tho' in region_input or any(k in full_text for k in can_tho_kw):
                region_code = 'can_tho'
                region_name = 'Cần Thơ'
                target_shop_names = ['Thế Giới Sơ Sinh', 'Nous Baby', 'Nhà Của Gạo']
            elif 'ho_chi_minh' in region_input or 'hcm' in region_input or any(k in full_text for k in hcm_kw):
                region_code = 'ho_chi_minh'
                region_name = 'Hồ Chí Minh'
                target_shop_names = ['Voi Con']
            elif 'ha_noi' in region_input or 'hn' in region_input or any(k in full_text for k in hn_kw):
                region_code = 'ha_noi'
                region_name = 'Hà Nội'
                target_shop_names = ['Mẹ Ơi', 'Sữa 3 Tốt']
            elif 'hai_phong' in region_input or 'hp' in region_input or any(k in full_text for k in hp_kw):
                region_code = 'hai_phong'
                region_name = 'Hải Phòng'
                target_shop_names = ['Ababy', 'Cute Baby']
            else:
                region_code = 'other'
                region_name = 'Khu vực khác (Toàn quốc)'
                target_shop_names = ['Long Châu']

            # Lấy danh sách shop object tương ứng
            available_shops = [s for s in shops if s.get('name') in target_shop_names]
            if not available_shops:
                available_shops = [s for s in shops if s.get('name') == 'Long Châu'] or shops[:1]

            # Chọn shop ưu tiên
            chosen_shop = None
            if pref_shop:
                for s in available_shops:
                    if s.get('name').lower() == pref_shop.lower():
                        chosen_shop = s
                        break

            if not chosen_shop:
                # Ưu tiên shop còn nhiều hạn mức khả dụng nhất / ít đơn chờ nhất
                chosen_shop = sorted(available_shops, key=lambda x: (x.get('pending_orders', 0), -x.get('remaining_quota', 0)))[0]

            code = body.get('code', f"Z{len(db.get('orders', [])) + 1:03d}")
            acc_num = chosen_shop.get('bank_account', '1057169817')
            bank_id = chosen_shop.get('bank_name', 'vietcombank').lower().replace(' ', '')
            # Chuẩn hóa bank_id cho VietQR
            if 'vietcom' in bank_id: bank_id = 'vietcombank'
            elif 'techcom' in bank_id: bank_id = 'techcombank'
            elif 'vp' in bank_id: bank_id = 'vpbank'
            elif 'bidv' in bank_id: bank_id = 'bidv'
            elif 'mb' in bank_id: bank_id = 'mbbank'
            elif 'vietin' in bank_id: bank_id = 'vietinbank'

            acc_name = chosen_shop.get('account_holder', 'HoKinhDoanhLaDucThanh')
            memo = f"{code} Tien sua Danone Zmoms"

            vietqr_url = f"https://img.vietqr.io/image/{bank_id}-{acc_num}-compact2.jpg?amount={amount}&addInfo={urllib.parse.quote(memo)}&accountName={urllib.parse.quote(acc_name)}"

            self.send_json({
                'success': True,
                'region_code': region_code,
                'region_name': region_name,
                'recommended_shop': chosen_shop.get('name'),
                'target_shop_names': target_shop_names,
                'available_shops': available_shops,
                'shop_details': chosen_shop,
                'vietqr_url': vietqr_url,
                'memo': memo,
                'amount': amount
            })

        elif path == '/api/orders':
            orders = db.get('orders', [])
            new_id = len(orders) + 1
            code = f"Z{new_id:03d}"
            
            now_str = datetime.datetime.now().strftime('%d/%m/%Y')
            assigned_shop_name = body.get('shop', 'Voi Con')

            new_order = {
                'id': new_id,
                'code': code,
                'date': now_str,
                'customer_name': body.get('customer_name', ''),
                'phone': body.get('phone', ''),
                'funnel': body.get('funnel', 'Thành viên Zmoms'),
                'address': body.get('address', ''),
                'district': body.get('district', ''),
                'product': body.get('product', 'Combo Aptamil 380g'),
                'shop': assigned_shop_name,
                'bill_link': body.get('bill_link', ''),
                'shop_status': '[1] Chờ tiếp nhận',
                'carrier': body.get('carrier', 'Ahamove'),
                'tracking_code': '',
                'pod_link': '',
                'overall_status': '[1] Chờ tiếp nhận',
                'source_channel': body.get('source_channel', 'Hội thảo BV'),
                'created_at': now_str
            }
            orders.insert(0, new_order)
            db['orders'] = orders

            # Cập nhật metrics cho Shop MBI được chọn
            shops = db.get('mbi_shops', [])
            for s in shops:
                if s.get('name') == assigned_shop_name:
                    s['total_orders'] = s.get('total_orders', 0) + 1
                    s['assigned'] = s.get('assigned', 0) + 1
                    s['pending_orders'] = s.get('pending_orders', 0) + 1
                    if s.get('remaining_quota', 0) > 0:
                        s['remaining_quota'] -= 1
                    quota = s.get('quota', 1)
                    s['completion_rate'] = round((s['total_orders'] / quota) * 100, 1) if quota else 0
                    break
            db['mbi_shops'] = shops

            # Thêm hoặc cập nhật Khách hàng
            customers = db.get('customers', [])
            cust = next((c for c in customers if c.get('phone') == new_order['phone']), None)
            if cust:
                cust['order_count'] = cust.get('order_count', 0) + 1
                if 'orders' not in cust: cust['orders'] = []
                cust['orders'].insert(0, code)
                cust['assigned_shop'] = assigned_shop_name
            else:
                customers.insert(0, {
                    'id': f"MOM-{len(customers)+1:04d}",
                    'name': new_order['customer_name'],
                    'phone': new_order['phone'],
                    'address': new_order['address'],
                    'district': new_order['district'],
                    'channel': new_order['source_channel'],
                    'funnel': new_order['funnel'],
                    'stage': '1. Mang thai 3 tháng đầu (Tuần 1 - 12)',
                    'codata_zmoms': True,
                    'codata_aptaclub': True,
                    'order_count': 1,
                    'orders': [code],
                    'assigned_shop': new_order['shop'],
                    'created_at': now_str
                })
            db['customers'] = customers

            # Lưu vào JSON và JS
            save_db(db)

            # Đồng bộ ghi dòng mới vào sheet ORDERS của file Excel
            append_order_to_excel(new_order)

            self.send_json({'success': True, 'order': new_order, 'message': f'✅ Đã tạo thành công đơn hàng {code}! Shop {assigned_shop_name} đã nhận đơn.'})

        elif path == '/api/customers':
            customers = db.get('customers', [])
            phone = str(body.get('phone', '')).replace(' ', '').strip()
            
            # Check duplicate phone
            existing = next((c for c in customers if c.get('phone') == phone and phone), None)
            if existing:
                self.send_json({'success': False, 'message': f'Số điện thoại {phone} đã tồn tại trên hệ thống ({existing.get("name")})!'}, 400)
                return

            now_str = datetime.datetime.now().strftime('%d/%m/%Y')
            new_cust = {
                'id': f"MOM-{len(customers)+1:04d}",
                'name': body.get('name', 'Mẹ Bầu'),
                'phone': phone,
                'address': body.get('address', ''),
                'district': body.get('district', 'TP.HCM'),
                'channel': body.get('channel', 'Hội thảo BV'),
                'funnel': body.get('funnel', 'Thành viên Zmoms'),
                'stage': body.get('stage', '2. Mang thai 3 tháng giữa (Tuần 13 - 27)'),
                'codata_zmoms': body.get('codata_zmoms', True),
                'codata_aptaclub': body.get('codata_aptaclub', True),
                'order_count': 0,
                'orders': [],
                'assigned_shop': body.get('assigned_shop', 'Voi Con'),
                'created_at': now_str
            }
            customers.insert(0, new_cust)
            db['customers'] = customers
            save_db(db)
            self.send_json({'success': True, 'customer': new_cust, 'message': 'Đã thêm mẹ bầu mới thành công!'})

        elif path == '/api/settings/update-targets':
            meta = db.get('meta', {})
            if 'target_data' in body: meta['target_data'] = int(body['target_data'])
            if 'target_nu' in body: meta['target_nu'] = int(body['target_nu'])
            if 'target_cr' in body: meta['target_cr'] = float(body['target_cr'])
            if 'project_name' in body: meta['project_name'] = str(body['project_name'])
            if 'duration' in body: meta['duration'] = str(body['duration'])
            if 'calc_mode' in body: meta['calc_mode'] = str(body['calc_mode'])
            if 'total_raw_data' in body and body['total_raw_data'] != '': meta['total_raw_data'] = int(body['total_raw_data'])
            if 'valid_raw_data' in body and body['valid_raw_data'] != '': meta['valid_raw_data'] = int(body['valid_raw_data'])
            if 'converted_raw_data' in body and body['converted_raw_data'] != '': meta['converted_raw_data'] = int(body['converted_raw_data'])
            if 'approved_hospitals' in body: meta['approved_hospitals'] = int(body['approved_hospitals'])
            if 'total_hospitals' in body: meta['total_hospitals'] = int(body['total_hospitals'])
            
            # Cập nhật Quota các shop MBI nếu có
            shop_quotas = body.get('shop_quotas', {})
            if shop_quotas:
                shops = db.get('mbi_shops', [])
                for s in shops:
                    s_name = s.get('name')
                    if s_name in shop_quotas:
                        try:
                            new_q = int(shop_quotas[s_name])
                            s['quota'] = new_q
                            s['remaining_quota'] = max(0, new_q - s.get('total_orders', 0))
                            s['completion_rate'] = round((s.get('total_orders', 0) / new_q) * 100, 1) if new_q > 0 else 0
                        except Exception:
                            pass
                db['mbi_shops'] = shops

            db['meta'] = meta
            save_db(db)
            self.send_json({'success': True, 'meta': meta, 'mbi_shops': db.get('mbi_shops', []), 'message': 'Đã cập nhật thông số mục tiêu chiến dịch thành công!'})

        elif path == '/api/kpi/entry':
            # Live Data Entry from team (matching Render)
            pic_code = body.get('pic_code', '')
            actual_nu = int(body.get('actual_nu', 0))
            subsidized_nu = int(body.get('subsidized_nu', 0))

            sales_team = db.get('sales_team', [])
            member = next((m for m in sales_team if m.get('code') == pic_code), None)
            if member:
                member['actual_nu'] = actual_nu
                member['subsidized_nu'] = subsidized_nu
                member['base_revenue'] = actual_nu * 750000
                save_db(db)
                self.send_json({'success': True, 'member': member, 'message': f'Đã cập nhật số thực đạt cho {member.get("name")} ({pic_code})!'})
            else:
                self.send_json({'success': False, 'message': 'Không tìm thấy nhân sự'}, 404)

        elif path == '/api/action-items':
            items = db.get('action_items', [])
            new_item = {
                'id': len(items) + 1,
                'task': body.get('task', ''),
                'owner': body.get('owner', 'Zmoms'),
                'pic': body.get('pic', ''),
                'deadline': body.get('deadline', ''),
                'status': body.get('status', '🟡 Đang xử lý')
            }
            items.insert(0, new_item)
            db['action_items'] = items
            save_db(db)
            self.send_json({'success': True, 'item': new_item})

        else:
            self.send_error(404, "Endpoint not found")

    def handle_api_put(self, path, body):
        db = load_db()

        if path.startswith('/api/orders/') and path.endswith('/status'):
            order_id = path.split('/')[3]
            orders = db.get('orders', [])
            target = next((o for o in orders if str(o.get('id')) == str(order_id) or str(o.get('code')) == str(order_id)), None)

            if target:
                if 'shop_status' in body: target['shop_status'] = body['shop_status']
                if 'carrier' in body: target['carrier'] = body['carrier']
                if 'tracking_code' in body: target['tracking_code'] = body['tracking_code']
                if 'pod_link' in body: target['pod_link'] = body['pod_link']
                if 'overall_status' in body: target['overall_status'] = body['overall_status']
                if 'notes' in body: target['notes'] = body['notes']

                save_db(db)
                self.send_json({'success': True, 'order': target, 'message': 'Cập nhật tiến độ đơn hàng thành công!'})
            else:
                self.send_json({'success': False, 'message': 'Không tìm thấy đơn hàng'}, 404)

        elif path.startswith('/api/weeks/'):
            week_name = path.split('/')[3] # e.g. W01, W02
            weeks = db.get('weeks', [])
            target = next((w for w in weeks if w.get('week') == week_name), None)
            if target:
                if 'target_data' in body: target['target_data'] = float(body['target_data'])
                if 'actual_data' in body: target['actual_data'] = float(body['actual_data'])
                if 'target_nu' in body: target['target_nu'] = float(body['target_nu'])
                if 'actual_nu' in body: target['actual_nu'] = float(body['actual_nu'])
                save_db(db)
                self.send_json({'success': True, 'week': target, 'message': f'Đã cập nhật chỉ số {week_name} thành công!'})
            else:
                self.send_json({'success': False, 'message': 'Không tìm thấy tuần'}, 404)

        elif path.startswith('/api/hospitals/'):
            hosp_id = path.split('/')[3]
            hospitals = db.get('hospitals', [])
            target = next((h for h in hospitals if str(h.get('id')) == str(hosp_id)), None)
            if target:
                if 'status' in body: target['status'] = body['status']
                if 'meeting_status' in body: target['meeting_status'] = body['meeting_status']
                save_db(db)
                self.send_json({'success': True, 'hospital': target})
            else:
                self.send_json({'success': False, 'message': 'Không tìm thấy bệnh viện'}, 404)

        elif path.startswith('/api/action-items/'):
            item_id = path.split('/')[3]
            items = db.get('action_items', [])
            target = next((i for i in items if str(i.get('id')) == str(item_id)), None)
            if target:
                if 'status' in body: target['status'] = body['status']
                save_db(db)
                self.send_json({'success': True, 'item': target})
            else:
                self.send_json({'success': False, 'message': 'Không tìm thấy task'}, 404)

        else:
            self.send_error(404, "Endpoint not found")

def run(port=PORT):
    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("0.0.0.0", port), CRMHandler) as httpd:
        print(f"======================================================")
        print(f"🚀 DANONE × ZMOMS OMNI-CHANNEL CRM SERVER RUNNING")
        print(f"👉 Local URL: http://localhost:{port}")
        print(f"======================================================")
        httpd.serve_forever()

if __name__ == '__main__':
    run()
